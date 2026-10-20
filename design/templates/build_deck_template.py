"""Build the Answer ALS deck template from design/tokens.json: python build_deck_template.py. Writes answer-als-deck-template.pptx to <data root>/answer-als/agent-handbook/v1/templates/. One example slide per layout; each slide's text says what goes where and its speaker notes hold the layout rule. Rerunnable; overwrites the deck."""
import json
import os
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    if sys.platform == "darwin" and not os.environ.get("HANDBOOK_ARM64"):
        os.execvpe("arch", ["arch", "-arm64", sys.executable, *sys.argv], dict(os.environ, HANDBOOK_ARM64="1"))
    raise
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
DESIGN = HERE.parent
sys.path.insert(0, str(HERE.parents[3]))
import board_check

VERSION = "answer-als/agent-handbook/v1/templates"
DECK_NAME = "answer-als-deck-template.pptx"
LAYOUTS = ["Title", "Section divider", "Headline plus text", "Two column", "Headline plus image", "Table", "Chart", "Numbers", "Quote or takeaway", "Closing"]
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
CHART_KEYS = ["chart_1", "chart_2", "chart_3", "chart_4", "chart_5", "chart_6"]


class Deck:
    def __init__(self, tokens):
        self.t = tokens
        self.c = tokens["color"]
        self.f = tokens["font"]
        self.pt = tokens["size"]["slide_pt"]
        self.size = tokens["size"]["slide_in"]
        self.logo = DESIGN / tokens["logo"]
        self.prs = Presentation()
        self.prs.slide_height = Inches(self.size["height"])
        self.prs.slide_width = Emu(self.prs.slide_height * 16 // 9)
        self.number = 0

    def rgb(self, name):
        return RGBColor.from_string(self.c[name])

    def new(self, background):
        self.number += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb(background)
        return slide

    def text(self, slide, left, top, width, height, lines, size, color, bold=False, spacing=None, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, gap=0, font=None):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.alignment = align
            para.line_spacing = 1.05
            if i:
                para.space_before = Pt(gap)
            parts = [(line, bold)] if isinstance(line, str) else [line] if isinstance(line, tuple) else line
            for part, part_bold in parts:
                run = para.add_run()
                run.text = part
                run.font.size = Pt(size)
                run.font.bold = part_bold
                run.font.name = font or self.f["heading" if part_bold else "body"]
                run.font.color.rgb = self.rgb(color)
                if spacing is not None:
                    run.font._element.set("spc", str(spacing))
        return box

    def bar(self, slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def notes(self, slide, rule):
        slide.notes_slide.notes_text_frame.text = rule

    def eyebrow(self, slide, left, top, width, text, color):
        return self.text(slide, left, top, width, 0.4, text, self.pt["eyebrow"], color, bold=True, spacing=self.f["eyebrow_tracking_slide_pt"] * 100)

    def brand(self, eyebrow, headline, subline, logo=True):
        slide = self.new("brand")
        if logo:
            slide.shapes.add_picture(str(self.logo), Inches(0.9), Inches(0.75), height=Inches(self.size["logo_height"]))
            self.text(slide, 1.95, 0.75, 10.4, 0.75, eyebrow, self.pt["eyebrow"], "tint_3", bold=True, spacing=self.f["eyebrow_tracking_slide_pt"] * 100, anchor=MSO_ANCHOR.MIDDLE)
        self.bar(slide, 0.9, 2.6, 0.9, 0.07, "tint_3")
        self.text(slide, 0.9, 2.9, 11.5, 2.4, headline, self.pt["title_headline"], "surface", bold=True)
        if subline:
            self.text(slide, 0.9, 5.95, 11.5, 0.5, subline, self.pt["subline"], "tint_3")
        return slide

    def content(self, eyebrow, headline, width=11.5, top=0.7, height=1.3):
        slide = self.new("surface")
        self.bar(slide, 0, 0, self.size["edge_strip"], self.size["height"], "brand")
        self.eyebrow(slide, 0.9, top, width, eyebrow, "brand")
        self.bar(slide, 0.9, top + 0.55, 0.9, 0.07, "brand")
        if headline:
            self.text(slide, 0.9, top + 0.85, width, height, headline, self.pt["headline"], "ink", bold=True)
        self.text(slide, 11.4, 6.85, 1.2, 0.35, str(self.number), self.pt["page_number"], "placeholder", align=PP_ALIGN.RIGHT)
        return slide

    def caption(self, slide, left, top, width, text):
        return self.text(slide, left, top, width, 0.4, text, self.pt["eyebrow"], "caption")

    def border(self, cell, side, color, width=Pt(1)):
        tcPr = cell._tc.get_or_add_tcPr()
        line = OxmlElement(f"a:{side}")
        line.set("w", str(int(width)))
        solid = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", self.c[color])
        solid.append(clr)
        line.append(solid)
        tcPr.append(line)


def title(d):
    s = d.brand([[("Program ", True), ("and event, 2 to 4 words", False)]],"Title: the point of the talk in under 10 words", "Presenter name, role and date")
    d.notes(s, "Title slide: brand blue, logo top left, one headline under 10 words and the presenter line, nothing else.")


def section(d):
    s = d.brand(None, "Section divider: name the next part in under 8 words", None, logo=False)
    d.notes(s, "Section divider: one headline on brand blue to mark a new part of the talk, no logo and no other text.")


def headline_text(d):
    s = d.content("TOPIC, 2 TO 4 WORDS", "Headline states the point in under 10 words")
    lines = ["Up to four short lines, 8 words or fewer each", "Each line backs up the headline", "Plain sentences, no bullets or sub-points", "Cut a line before you shrink the text"]
    d.text(s, 0.9, 3.3, 11.0, 3.2, lines, d.pt["subline"], "text", gap=14)
    d.notes(s, "Headline plus text: one point and up to four short lines that support it; if you need more lines, split the slide.")


def two_column(d):
    s = d.content("COMPARISON, 2 TO 4 WORDS", "Two columns compare two things side by side")
    left = [("Left label, 1 to 3 words", True), "Up to three short lines", "Same order of points as the right", "8 words or fewer per line"]
    right = [("Right label, 1 to 3 words", True), "Up to three short lines", "Match each line to the left", "Keep both columns the same length"]
    d.text(s, 0.9, 3.3, 5.4, 3.2, left, d.pt["subline"], "text", gap=14)
    d.text(s, 6.9, 3.3, 5.4, 3.2, right, d.pt["subline"], "text", gap=14)
    d.notes(s, "Two column: compare two things, before and after or this and that, with the same number of lines on each side.")


def headline_image(d):
    s = d.content("TOPIC, 2 TO 4 WORDS", "Headline left says what the picture shows", width=5.0, top=1.7, height=2.6)
    d.text(s, 0.9, 5.2, 5.0, 1.0, "One short line on what to look at, 12 words at most", d.pt["subline"], "text")
    slot = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.7), Inches(0.9), Inches(5.9), Inches(5.0))
    slot.name = "Picture slot"
    slot.fill.solid()
    slot.fill.fore_color.rgb = d.rgb("tint_1")
    slot.line.color.rgb = d.rgb("rule")
    slot.line.width = Pt(0.75)
    slot.shadow.inherit = False
    frame = slot.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(["Picture slot", "Delete this box and insert one screenshot or photo here.", "Keep its sides locked and frame it with a thin rule line."]):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = line
        run.font.size = Pt(d.pt["command"])
        run.font.bold = i == 0
        run.font.name = d.f["body"]
        run.font.color.rgb = d.rgb("placeholder")
    d.caption(s, 6.7, 6.05, 5.9, "Caption: what the picture is and where it came from, one sentence")
    d.notes(s, "Headline plus image: one screenshot or photo on the right with a caption, and the headline on the left says why it matters.")


def table(d):
    s = d.content("TABLE, 2 TO 4 WORDS", "Headline says what the table shows")
    rows = [
        ["Column", "What goes here", "Max words", "Align"],
        ["Label", "Plain words, never a raw column name", "3", "Left"],
        ["Count", "Numbers in one unit per column", "1", "Right"],
        ["Release", "Release name, such as Release 8", "2", "Left"],
        ["Note", "Short phrase or leave the cell empty", "6", "Left"],
    ]
    numeric = {2}
    frame = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.9), Inches(2.9), Inches(11.5), Inches(0.6 * len(rows)))
    tbl = frame.table
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = NO_STYLE_TABLE
    tbl.first_row = True
    tbl.horz_banding = False
    for c, width in enumerate([2.3, 5.6, 1.8, 1.8]):
        tbl.columns[c].width = Inches(width)
    for r, values in enumerate(rows):
        tbl.rows[r].height = Inches(0.6)
        for c, value in enumerate(values):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.RIGHT if c in numeric and r else PP_ALIGN.LEFT
            run = para.add_run()
            run.text = value
            run.font.size = Pt(d.pt["command"])
            run.font.bold = r == 0
            run.font.name = d.f["heading" if r == 0 else "body"]
            run.font.color.rgb = d.rgb("ink" if r == 0 else "text_2")
            d.border(cell, "lnL", "table_border")
            d.border(cell, "lnR", "table_border")
            d.border(cell, "lnT", "table_head_rule" if r == 1 else "table_border")
            d.border(cell, "lnB", "table_head_rule" if r == 0 else "table_border")
            cell.fill.solid()
            cell.fill.fore_color.rgb = d.rgb("tint_1" if r == 0 else "surface")
    d.caption(s, 0.9, 6.15, 11.5, "Source: name the release and file the table comes from")
    d.notes(s, "Table: at most 5 rows and 4 columns with plain word headers, numbers right aligned, and a source line under it.")


def chart(d):
    s = d.content("CHART, 2 TO 4 WORDS", "Headline states what the chart shows")
    data = CategoryChartData()
    data.categories = ["Category 1", "Category 2", "Category 3", "Category 4"]
    for name, values in [("Series 1", (40, 55, 30, 62)), ("Series 2", (35, 48, 42, 50)), ("Series 3", (20, 26, 38, 33))]:
        data.add_series(name, values)
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(2.75), Inches(11.5), Inches(3.55), data)
    ch = frame.chart
    ch.has_title = False
    ch.font.name = d.f["body"]
    ch.font.size = Pt(d.pt["command"])
    ch.font.color.rgb = d.rgb("chart_axis")
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    plot = ch.plots[0]
    plot.gap_width = 80
    plot.overlap = 0
    plot.vary_by_categories = False
    for series, key in zip(plot.series, CHART_KEYS):
        series.invert_if_negative = False
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = d.rgb(key)
        series.format.line.fill.background()
    for axis, label in [(ch.category_axis, "Category axis label in sentence case"), (ch.value_axis, "Value with its unit")]:
        axis.has_title = True
        axis.axis_title.text_frame.text = label
        run = axis.axis_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(d.pt["command"])
        run.font.bold = False
        run.font.name = d.f["body"]
        run.font.color.rgb = d.rgb("chart_axis")
        axis.tick_labels.font.size = Pt(d.pt["command"])
        axis.tick_labels.font.color.rgb = d.rgb("chart_axis")
        axis.format.line.color.rgb = d.rgb("chart_grid")
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = d.rgb("chart_grid")
    ch.value_axis.major_gridlines.format.line.width = Pt(0.75)
    ch.category_axis.has_major_gridlines = False
    d.caption(s, 0.9, 6.35, 11.5, "Source: name the release, such as Answer ALS Release 7 index")
    d.notes(s, "Chart: one native chart in the palette colors in order, at most six series, every axis labeled with its unit and a source line under it.")


def numbers(d):
    s = d.content("BY THE NUMBERS, 2 TO 4 WORDS", "Headline says what the three numbers mean")
    for i, (value, label) in enumerate([("000", "Label, 1 to 3 words"), ("00%", "Label, 1 to 3 words"), ("0.0", "Label, 1 to 3 words")]):
        left = 0.9 + i * 3.9
        d.text(s, left, 3.3, 3.6, 1.0, value, d.pt["title_headline"], "brand", bold=True)
        d.text(s, left, 4.45, 3.6, 0.8, label, d.pt["subline"], "text")
    d.caption(s, 0.9, 6.15, 11.5, "Source: published figures only, name where each was published")
    d.notes(s, "Numbers: three published figures at most, each with a short label, and a source line naming where they were published.")


def takeaway(d):
    s = d.content("TAKEAWAY", None)
    d.text(s, 0.9, 2.2, 11.5, 2.8, "One sentence the audience should remember, under 15 words", d.pt["title_headline"], "brand", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, 0.9, 5.4, 11.5, 0.5, "Name and role for a quote, or delete this line for a takeaway", d.pt["subline"], "caption")
    d.notes(s, "Quote or takeaway: one sentence set large with nothing else on the slide but its attribution.")


def closing(d):
    s = d.brand("Your challenge:", "Closing: one action the audience takes next, under 10 words", "Presenter name and how to reach the team")
    d.notes(s, "Closing slide: brand blue, logo top left, one action for the audience and the presenter line, the mirror of the title slide.")


BUILDERS = [title, section, headline_text, two_column, headline_image, table, chart, numbers, takeaway, closing]


def build(path):
    tokens = json.loads((DESIGN / "tokens.json").read_text(encoding="ascii"))
    d = Deck(tokens)
    for make in BUILDERS:
        make(d)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.prs.core_properties.author = d.prs.core_properties.last_modified_by = "Answer ALS data team"
    d.prs.save(path)
    return path


if __name__ == "__main__":
    deck = build(board_check.data_path("") / VERSION / DECK_NAME)
    print(f"wrote {deck.as_posix()}: {len(BUILDERS)} slides")
