"""Build the Answer ALS Word templates from design/tokens.json: python build_docx_templates.py. Writes answer-als-memo-template.docx and answer-als-one-pager-template.docx to <data root>/answer-als/agent-handbook/v1/templates/. Letter, 1 inch margins, styles named Title, Heading 1, Heading 2, Body, Caption, Callout and Table Grid; every paragraph's text is its own instruction. Rerunnable; overwrites both files with the same bytes each time."""
import datetime
import json
import os
import re
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
except ImportError:
    if sys.platform == "darwin" and not os.environ.get("HANDBOOK_ARM64"):
        os.execvpe("arch", ["arch", "-arm64", sys.executable, *sys.argv], dict(os.environ, HANDBOOK_ARM64="1"))
    raise
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DESIGN = HERE.parent
sys.path.insert(0, str(HERE.parents[3]))
import board_check

VERSION = "answer-als/agent-handbook/v1/templates"
MEMO_NAME = "answer-als-memo-template.docx"
ONE_PAGER_NAME = "answer-als-one-pager-template.docx"
STYLES = ["Title", "Heading 1", "Heading 2", "Body", "Caption", "Callout", "Table Grid"]
PROGRAM = "Answer ALS"
STAMP = datetime.datetime(2026, 10, 20, 0, 0, 0)
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def order(ppr):
    children = sorted(ppr, key=lambda e: PPR_ORDER.index(e.tag.split("}")[1]))
    for child in children:
        ppr.remove(child)
        ppr.append(child)


class Doc:
    def __init__(self, tokens, title):
        self.c = tokens["color"]
        self.f = tokens["font"]
        self.pt = tokens["size"]["document_pt"]
        self.line = Pt(self.pt["body"] * self.f["line_height"])
        self.doc = Document()
        self.page()
        self.fonts()
        self.styles()
        self.footer()
        props = self.doc.core_properties
        props.title = title
        props.author = props.last_modified_by = "Answer ALS data team"
        props.comments = "Built by answer-als/agent-handbook/design/templates/build_docx_templates.py from design/tokens.json."
        props.created = props.modified = STAMP
        props.revision = 1
        props.keywords = props.subject = props.category = ""

    def page(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.5)

    def fonts(self):
        body, fallback = self.f["body"], self.f["fallback"]
        part = next(p for p in self.doc.part.package.iter_parts() if p.partname == "/word/fontTable.xml")
        xml = part.blob.decode("utf-8")
        entry = f'<w:font w:name="{body}"><w:altName w:val="{fallback}"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>'
        part._blob = re.sub(r"(<w:fonts[^>]*>)", lambda m: m.group(1) + entry, xml, count=1).encode("utf-8")
        defaults = self.doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        old = defaults.find(qn("w:rFonts"))
        if old is not None:
            defaults.remove(old)
        defaults.insert(0, parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{body}" w:hAnsi="{body}" w:cs="{fallback}" w:eastAsia="{body}"/>'))

    def rgb(self, name):
        return RGBColor.from_string(self.c[name])

    def reset(self, style):
        for tag in ("w:pPr", "w:rPr", "w:link"):
            old = style.element.find(qn(tag))
            if old is not None:
                style.element.remove(old)

    def para_style(self, name, size, color, bold=False, before=0, after=6, line=None, base="Normal", next_name="Body"):
        styles = self.doc.styles
        style = styles[name] if name in [s.name for s in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        self.reset(style)
        style.base_style = styles[base] if base else None
        style.next_paragraph_style = styles[next_name] if next_name in [s.name for s in styles] else style
        style.quick_style = True
        font = style.font
        font.name = self.f["heading" if bold else "body"]
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = self.rgb(color)
        fmt = style.paragraph_format
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        if line is None:
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        else:
            fmt.line_spacing = line
        return style

    def styles(self):
        self.para_style("Normal", self.pt["body"], "text", after=0, base=None, next_name="Normal")
        self.para_style("Body", self.pt["body"], "text", after=6, line=self.line)
        title = self.para_style("Title", self.pt["title"], "brand", bold=True, after=6)
        h1 = self.para_style("Heading 1", self.pt["heading"], "brand", bold=True, before=16, after=6)
        h2 = self.para_style("Heading 2", self.pt["body"], "brand", bold=True, before=12, after=4)
        for style in (title, h1, h2):
            style.paragraph_format.keep_with_next = True
        for level, style in enumerate((h1, h2)):
            style.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level}"/>'))
            order(style.element.pPr)
        self.para_style("Caption", self.pt["caption"], "caption", after=6)
        callout = self.para_style("Callout", self.pt["body"], "text", after=0, line=self.line)
        tint = self.c["tint_1"]
        sides = "".join(f'<w:{side} w:val="single" w:sz="4" w:space="6" w:color="{tint}"/>' for side in ("top", "left", "bottom", "right"))
        ppr = callout.element.get_or_add_pPr()
        ppr.insert(0, parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{tint}"/>'))
        ppr.insert(0, parse_xml(f'<w:pBdr {nsdecls("w")}>{sides}</w:pBdr>'))
        ppr.insert(0, parse_xml(f'<w:ind {nsdecls("w")} w:left="120" w:right="120"/>'))
        order(ppr)
        grid = self.doc.styles["Table Grid"]
        tblpr = grid.element.find(qn("w:tblPr"))
        tblpr.remove(tblpr.find(qn("w:tblBorders")))
        line = self.c["table_border"]
        edges = "".join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="{line}"/>' for side in ("top", "left", "bottom", "right", "insideH", "insideV"))
        tblpr.insert(1, parse_xml(f'<w:tblBorders {nsdecls("w")}>{edges}</w:tblBorders>'))
        margins = tblpr.find(qn("w:tblCellMar"))
        for side in ("top", "bottom"):
            margins.find(qn(f"w:{side}")).set(qn("w:w"), "60")
        grid.font.name = self.f["body"]
        grid.font.size = Pt(self.pt["body"])
        grid.font.color.rgb = self.rgb("text_2")

    def footer(self):
        footer = self.doc.sections[0].footer
        para = footer.paragraphs[0]
        para.style = self.doc.styles["Caption"]
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        self.run(para, f"{PROGRAM}\tPage ", color="placeholder")
        field = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
        para._p.append(field)
        for run in para.runs:
            run.font.color.rgb = self.rgb("placeholder")
        field.find(qn("w:r")).insert(0, parse_xml(f'<w:rPr {nsdecls("w")}><w:color w:val="{self.c["placeholder"]}"/></w:rPr>'))

    def run(self, para, text, bold=None, size=None, color=None):
        run = para.add_run(text)
        if bold is not None:
            run.bold = bold
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = self.rgb(color)
        return run

    def add(self, style, text="", **run):
        para = self.doc.add_paragraph(style=style)
        if text:
            self.run(para, text, **run)
        return para

    def rule(self, para):
        ppr = para._p.get_or_add_pPr()
        ppr.insert(0, parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="6" w:color="{self.c["brand"]}"/></w:pBdr>'))
        order(ppr)
        para.paragraph_format.space_after = Pt(16)

    def table(self, rows, widths, head=True):
        table = self.doc.add_table(rows=len(rows), cols=len(widths), style=self.doc.styles["Table Grid"])
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for r, values in enumerate(rows):
            for i, value in enumerate(values):
                cell = table.cell(r, i)
                cell.width = Inches(widths[i])
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Body"]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                self.run(para, value, bold=True if head and r == 0 else None, color="ink" if head and r == 0 else "text_2")
                if head and r == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    tcpr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="0" w:color="{self.c["table_head_rule"]}"/></w:tcBorders>'))
                    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{self.c["table_head"]}"/>'))
        return table

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        with zipfile.ZipFile(path) as z:
            items = [(info.filename, z.read(info)) for info in z.infolist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in items:
                z.writestr(zipfile.ZipInfo(name, STAMP.timetuple()[:6]), data, zipfile.ZIP_DEFLATED)
        return path


def tokens():
    return json.loads((DESIGN / "tokens.json").read_text(encoding="ascii"))


def memo(path):
    d = Doc(tokens(), "Answer ALS memo template")
    d.add("Title", "Memo")
    block = [
        ("To", "[Name and role of each reader. Roles only when the memo leaves the team.]"),
        ("From", "[Your name and role.]"),
        ("Date", "[Date sent, written as 20 October 2026.]"),
        ("Subject", "[One short sentence that states the point, under 12 words.]"),
    ]
    for i, (label, text) in enumerate(block):
        para = d.add("Body")
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.tab_stops.add_tab_stop(Inches(1))
        d.run(para, f"{label}:\t", bold=True, color="ink")
        d.run(para, text)
        if i == len(block) - 1:
            d.rule(para)
    sections = [
        ("Summary", ["[The point of the memo and what you need from the reader, in two or three sentences. Under 60 words. A reader who stops here should know the answer.]"]),
        ("Background", ["[Why this came up and what the reader needs to know first. Under 120 words. Name the release and the file, never a participant or a row.]"]),
        ("What we did", ["[What was done, in order, and what it showed. Under 200 words. Use a Heading 2 for each step when there are more than two.]"], "[Step name, in sentence case]", "[What this step did and what it found. Under 80 words. Every number comes from the release index or the published release notes, with the release named.]"),
        ("What is next", ["[The next steps, who owns each one by role, and when. Under 100 words. Say what is not planned as well.]"]),
    ]
    for heading, texts, *sub in sections:
        d.add("Heading 1", heading)
        for text in texts:
            d.add("Body", text)
        if sub:
            d.add("Heading 2", sub[0])
            d.add("Body", sub[1])
    d.add("Heading 1", "Ask")
    d.add("Callout", "[The one decision or action you need, from whom and by when. Under 40 words. Delete this section if there is no ask.]")
    d.add("Caption", "[Optional source line: Source: Answer ALS Release 7 index, or the file the numbers came from. Delete when the memo has no numbers.]").paragraph_format.space_before = Pt(12)
    return d.save(path)


def one_pager(path):
    d = Doc(tokens(), "Answer ALS one-pager template")
    d.rule(d.add("Title", "[Headline: one sentence that states the point, under 12 words]"))
    d.add("Body", "[One paragraph, under 90 words: what this is, who it is for and why it matters now. Plain words, no jargon, no participant identifiers. Keep the whole page to one page; cut words before you shrink the type.]")
    numbers = d.doc.add_table(rows=2, cols=3)
    numbers.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i in range(3):
        top = numbers.cell(0, i).paragraphs[0]
        top.style = d.doc.styles["Body"]
        top.paragraph_format.space_before = Pt(12)
        top.paragraph_format.space_after = Pt(0)
        top.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        d.run(top, f"[Number {i + 1}]", bold=True, size=d.pt["title"], color="brand")
        label = numbers.cell(1, i).paragraphs[0]
        label.style = d.doc.styles["Caption"]
        d.run(label, "[Label: what it counts. From the release index; name the release.]")
        for r in range(2):
            numbers.cell(r, i).width = Inches(6.5 / 3)
    d.add("Heading 1", "[Table heading: what the table shows]")
    d.table([
        ["[Column 1 header]", "[Column 2 header]", "[Column 3 header]"],
        ["[Row label]", "[Value]", "[Value]"],
        ["[Row label]", "[Value]", "[Value]"],
        ["[Row label]", "[Value]", "[Value]"],
    ], [2.5, 2.0, 2.0])
    d.add("Caption", "[Source line: the release and file the table comes from. Four rows at most; numbers right aligned, text left aligned.]").paragraph_format.space_before = Pt(4)
    d.add("Heading 1", "How to get access")
    d.add("Callout", "[Two or three sentences, under 50 words: where to apply, what the applicant needs, and how long approval takes. Link the portal page by name, not a SAS or expiring link.]")
    contact = d.add("Body")
    contact.paragraph_format.space_before = Pt(16)
    d.run(contact, "Contact: ", bold=True, color="ink")
    d.run(contact, "[Team role and shared team address, never a personal phone number.]")
    return d.save(path)


BUILDERS = {MEMO_NAME: memo, ONE_PAGER_NAME: one_pager}


def build(folder):
    return [make(Path(folder) / name) for name, make in BUILDERS.items()]


if __name__ == "__main__":
    for path in build(board_check.data_path("") / VERSION):
        print(f"wrote {path.as_posix()}")
